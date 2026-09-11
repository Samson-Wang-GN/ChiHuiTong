import io
import json
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta

from django.db import close_old_connections
from django.test import TestCase, TransactionTestCase
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from chihuitong.errors import BusinessError
from chihuitong.models import (
    Benefit,
    Card,
    Customer,
    CustomerSource,
)
from chihuitong.services import catalog, customers, files, imports, jobs, sales

from .support import actor_fixture, api_client
from .test_catalog import contract_fixture, image_bytes, product_fixture


def xlsx_bytes(rows, sheet="客户表"):
    book = Workbook()
    book.active.title = sheet
    for row in rows:
        book.active.append(row)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def sales_setup(test):
    test.platform = actor_fixture()
    test.resource = actor_fixture("broker", "13900000002")
    test.other = actor_fixture("insurance", "13900000003")
    test.staff = actor_fixture("broker", "13900000004", "staff", test.resource.organization)
    test.product = product_fixture(test.platform)
    for actor in [test.resource, test.other]:
        contract_fixture(test.platform, actor.organization, product=test.product)
    test.brand = catalog.save_brand(
        test.platform, test.resource.organization.id, name="合成保险福利"
    )
    test.other_brand = catalog.save_brand(
        test.platform, test.other.organization.id, name="另一方福利"
    )


def order_fixture(
    test, *, actor=None, brand=None, price=0, physical=False, quantity=2, customer=None
):
    return sales.create_order(
        actor or test.resource,
        product_id=test.product.id,
        source_brand_id=(brand or test.brand).id,
        mode="physical" if physical else "named",
        entry="quantity" if physical else "single",
        quantity=quantity if physical else None,
        rows=None
        if physical
        else [
            customer
            or {
                "name": "合成客户甲",
                "phone": "13900000101",
                "quantity": quantity,
                "resource_customer_no": "00001",
                "age": 0,
            }
        ],
        unit_price_cents=price,
    )


def approve(test, order):
    return sales.approve_order(
        test.platform, order.id, approved=True, version=order.version, reason="合成开卡审核"
    )


class SalesTests(TestCase):
    def setUp(self):
        sales_setup(self)

    def test_named_pending_no_customer_until_approved_and_claim_starts_expiry(self):
        order = order_fixture(self)
        self.assertEqual(Customer.objects.count(), 0)
        self.assertEqual(Card.objects.count(), 0)
        approve(self, order)
        self.assertEqual(Customer.objects.count(), 1)
        self.assertEqual(Benefit.objects.count(), 2)
        card = Card.objects.filter(order=order).first()
        benefit = card.benefit
        self.assertEqual((benefit.pending, benefit.available), (1, 0))
        self.assertIsNone(benefit.expires_at)
        customer = customers.register_verified_customer("13900000101")
        benefit = customers.claim(customer.id, card.id)
        self.assertEqual((benefit.pending, benefit.available), (0, 1))
        self.assertEqual(benefit.expires_at - benefit.activated_at, timedelta(days=180))
        self.assertEqual(customers.claim(customer.id, card.id).id, benefit.id)
        self.assertEqual(Benefit.objects.count(), 2)

    def test_shared_customer_external_number_not_identity_optional_only_fill_empty(self):
        first = approve(self, order_fixture(self))
        customer = Customer.objects.get()
        initial_id, initial_number = customer.id, customer.number
        self.assertNotEqual(initial_number, "00001")
        second = order_fixture(
            self,
            actor=self.other,
            brand=self.other_brand,
            customer={
                "name": "合成客户甲",
                "phone": "+86 139-0000-0101",
                "quantity": 1,
                "resource_customer_no": "NEW-0002",
                "age": 28,
                "gender": "女",
                "occupation": "教师",
            },
        )
        approve(self, second)
        customer.refresh_from_db()
        self.assertEqual(Customer.objects.count(), 1)
        self.assertEqual((customer.id, customer.number), (initial_id, initial_number))
        self.assertEqual(customer.profile["age"], 0)
        self.assertEqual(customer.profile["gender"], "female")
        self.assertEqual(CustomerSource.objects.count(), 2)
        third = order_fixture(
            self,
            customer={
                "name": "合成客户乙",
                "phone": "13900000102",
                "quantity": 1,
                "resource_customer_no": initial_number,
            },
        )
        approve(self, third)
        self.assertEqual(Customer.objects.count(), 2)
        self.assertEqual(first.cards.count(), 2)

    def test_name_conflict_blocks_order_and_approval_rechecks_atomically(self):
        pending = order_fixture(self)
        customers.register_verified_customer("13900000101", name="另一个合成姓名")
        with self.assertRaisesMessage(BusinessError, "手机号与姓名"):
            approve(self, pending)
        self.assertEqual(Card.objects.count(), 0)
        with self.assertRaises(BusinessError):
            order_fixture(self)

    def test_physical_no_excel_unique_range_activation_and_nontransfer(self):
        order = approve(self, order_fixture(self, physical=True, quantity=3))
        self.assertEqual(Customer.objects.count(), 0)
        cards = list(order.cards.order_by("serial"))
        self.assertEqual(order.number_range.last_number - order.number_range.first_number + 1, 3)
        customer = customers.register_verified_customer("13900000101", name="合成客户甲")
        other = customers.register_verified_customer("13900000102", name="合成客户乙")
        benefit = customers.activate(customer.id, cards[0].credential)
        self.assertEqual(customers.activate(customer.id, cards[0].credential).id, benefit.id)
        with self.assertRaises(BusinessError):
            customers.activate(other.id, cards[0].credential)
        second = approve(self, order_fixture(self, physical=True, quantity=2))
        self.assertGreater(second.number_range.first_number, order.number_range.last_number)

    def test_purchase_receipt_full_payment_before_approval(self):
        order = order_fixture(self, price=100)
        with self.assertRaises(BusinessError):
            approve(self, order)
        asset = files.upload_file(
            self.resource, data=image_bytes(), filename="合成回单.png", purpose="payment"
        )
        receipt = sales.submit_receipt(
            self.resource,
            order.id,
            amount_cents=100,
            paid_at=timezone.now(),
            payer="合成资源方",
            reference="SYNTH-ONE",
            attachment_ids=[str(asset.id)],
            version=order.version,
        )
        sales.review_receipt(
            self.platform,
            receipt.id,
            approved=True,
            version=receipt.version,
            reason="确认实际到账100分",
        )
        order.refresh_from_db()
        self.assertEqual(order.status, "pending_payment")
        with self.assertRaises(BusinessError):
            approve(self, order)
        receipt = sales.submit_receipt(
            self.resource,
            order.id,
            amount_cents=100,
            paid_at=timezone.now(),
            payer="合成资源方",
            reference="SYNTH-TWO",
            attachment_ids=[str(asset.id)],
            version=order.version,
        )
        sales.review_receipt(
            self.platform,
            receipt.id,
            approved=True,
            version=receipt.version,
            reason="确认剩余到账100分",
        )
        order.refresh_from_db()
        self.assertEqual(order.status, "pending_approval")
        approve(self, order)
        self.assertEqual(Card.objects.count(), 2)

    def test_receipt_rejection_and_duplicate_reference_no_double_money(self):
        order = order_fixture(self, price=100)
        asset = files.upload_file(
            self.resource, data=image_bytes(), filename="合成回单.png", purpose="payment"
        )

        def submit(ref):
            order.refresh_from_db()
            return sales.submit_receipt(
                self.resource,
                order.id,
                amount_cents=100,
                paid_at=timezone.now(),
                payer="合成资源方",
                reference=ref,
                attachment_ids=[str(asset.id)],
                version=order.version,
            )

        first = submit("SYNTH-RETRY")
        sales.review_receipt(
            self.platform, first.id, approved=False, version=first.version, reason="未查到"
        )
        second = submit("SYNTH-RETRY")
        sales.review_receipt(
            self.platform, second.id, approved=True, version=second.version, reason="已查到"
        )
        third = submit("SYNTH-RETRY")
        with self.assertRaises(BusinessError):
            sales.review_receipt(
                self.platform, third.id, approved=True, version=third.version, reason="重复流水"
            )
        order.refresh_from_db()
        self.assertEqual(order.received_cents, 100)

    def test_cancel_whole_batch_or_stop_only_unclaimed(self):
        order = approve(self, order_fixture(self))
        customer = customers.register_verified_customer("13900000101")
        cards = list(order.cards.order_by("serial"))
        customers.claim(customer.id, cards[0].id)
        with self.assertRaises(BusinessError):
            sales.request_stop(
                self.resource, order.id, action="cancel", version=order.version, reason="申请取消"
            )
        order = sales.request_stop(
            self.resource, order.id, action="stop", version=order.version, reason="停止剩余"
        )
        sales.review_stop(
            self.platform, order.id, approved=True, version=order.version, reason="确认停止"
        )
        self.assertEqual(Benefit.objects.get(card=cards[0]).available, 1)
        self.assertEqual(Benefit.objects.get(card=cards[1]).voided, 1)
        with self.assertRaises(BusinessError):
            customers.claim(customer.id, cards[1].id)

    def test_cancelled_ranges_never_reused(self):
        order = approve(self, order_fixture(self, physical=True))
        end = order.number_range.last_number
        order = sales.request_stop(
            self.resource, order.id, action="cancel", version=order.version, reason="整单取消"
        )
        sales.review_stop(
            self.platform, order.id, approved=True, version=order.version, reason="尚未激活"
        )
        self.assertEqual(order.cards.filter(status="void").count(), 2)
        next_order = approve(self, order_fixture(self, physical=True))
        self.assertGreater(next_order.number_range.first_number, end)

    def test_freeze_blocks_activation_and_cannot_change_owner(self):
        order = approve(self, order_fixture(self, physical=True))
        card = order.cards.first()
        customers.register_verified_customer("13900000101", name="合成客户")
        sales.freeze_card(self.platform, card.id, frozen=True, reason="争议核对")
        with self.assertRaises(BusinessError):
            customers.activate(Customer.objects.get().id, card.credential)
        with self.assertRaises(BusinessError):
            sales.freeze_card(self.resource, card.id, frozen=False, reason="越权")

    def test_staff_scope_and_idempotent_api(self):
        owner_order = order_fixture(self)
        own = order_fixture(self, actor=self.staff)
        self.assertEqual(sales.visible_orders(self.staff).count(), 1)
        client = api_client(self.staff)
        self.assertEqual(client.get(f"/api/v1/sales-orders/{owner_order.id}").status_code, 404)
        self.assertEqual(client.get(f"/api/v1/sales-orders/{own.id}").status_code, 200)
        payload = {
            "product_id": str(self.product.id),
            "source_brand_id": str(self.brand.id),
            "mode": "physical",
            "entry": "quantity",
            "quantity": 2,
        }
        response = client.post(
            "/api/v1/sales-orders",
            payload,
            format="json",
            HTTP_IDEMPOTENCY_KEY="SYNTH-REQUEST-0001",
        )
        repeated = client.post(
            "/api/v1/sales-orders",
            payload,
            format="json",
            HTTP_IDEMPOTENCY_KEY="SYNTH-REQUEST-0001",
        )
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()["id"], repeated.json()["id"])
        payload["quantity"] = 3
        self.assertEqual(
            client.post(
                "/api/v1/sales-orders",
                payload,
                format="json",
                HTTP_IDEMPOTENCY_KEY="SYNTH-REQUEST-0001",
            ).status_code,
            409,
        )

    def test_plain_export_has_no_executable_formulas(self):
        order = approve(self, order_fixture(self, physical=True))
        response = api_client(self.resource).post(
            f"/api/v1/sales-orders/{order.id}/cards/export", {"confirmed": True}, format="json"
        )
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.active.max_row, 3)
        self.assertEqual(book.active["A2"].data_type, "s")
        self.assertEqual(book.active["D2"].data_type, "s")


class ImportTests(TestCase):
    def setUp(self):
        sales_setup(self)

    def prepare(self, data, *, mapping=None, uniform=None):
        asset = files.upload_file(
            self.resource, data=xlsx_bytes(data), filename="合成客户.xlsx", purpose="sales_excel"
        )
        batch = imports.create_import(self.resource, asset.id)
        self.assertTrue(jobs.run_one())
        batch.refresh_from_db()
        batch = imports.configure_import(
            self.resource,
            batch.id,
            sheet="客户表",
            header_row=1,
            mapping=mapping or {"name": 0, "phone": 1, "quantity": 2, "resource_customer_no": 3},
            quantity_mode="uniform" if uniform else "column",
            uniform_quantity=uniform,
            version=batch.version,
        )
        self.assertTrue(jobs.run_one())
        batch.refresh_from_db()
        return batch

    def test_async_inspection_validation_confirmation_and_order(self):
        batch = self.prepare(
            [
                ["姓名", "手机号", "开卡数量", "客户编号"],
                ["合成客户甲", "13900000101", 2, "00012"],
                [None, None, None, None],
                ["合成客户乙", "13900000102", 1, "00013"],
            ]
        )
        self.assertEqual(batch.status, "validated")
        self.assertEqual(
            list(batch.rows.order_by("row_number").values_list("row_number", flat=True)), [2, 4]
        )
        self.assertEqual(Customer.objects.count(), 0)
        with self.assertRaises(BusinessError):
            sales.create_order(
                self.resource,
                product_id=self.product.id,
                source_brand_id=self.brand.id,
                mode="named",
                entry="excel",
                import_batch_id=batch.id,
            )
        batch = imports.confirm_import(
            self.resource, batch.id, mapping_digest=batch.mapping_digest, version=batch.version
        )
        order = sales.create_order(
            self.resource,
            product_id=self.product.id,
            source_brand_id=self.brand.id,
            mode="named",
            entry="excel",
            import_batch_id=batch.id,
        )
        approve(self, order)
        self.assertEqual(Card.objects.count(), 3)
        self.assertEqual(
            order.rows.order_by("row_number").first().normalized["resource_customer_no"], "00012"
        )
        with self.assertRaises(BusinessError):
            imports.configure_import(
                self.resource,
                batch.id,
                sheet="客户表",
                header_row=1,
                mapping={"name": 0, "phone": 1, "quantity": 2},
                quantity_mode="column",
                version=batch.version,
            )

    def test_errors_duplicates_formula_and_all_error_export(self):
        batch = self.prepare(
            [
                ["姓名", "手机号", "数量", "客户编号"],
                ["合成甲", "13900000101", 1, "01"],
                ["合成甲", "13900000101", 1, "02"],
                ["公式", "13900000102", "=1+1", "03"],
                ["缺数量", "13900000103", None, "04"],
            ]
        )
        self.assertEqual(batch.error_rows, 4)
        with self.assertRaises(BusinessError):
            imports.confirm_import(
                self.resource, batch.id, mapping_digest=batch.mapping_digest, version=batch.version
            )
        response = api_client(self.resource).get(
            f"/api/v1/imports/{batch.id}/errors.xlsx?page_size=1"
        )
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.active.max_row, 5)
        self.assertEqual(book.active["D4"].data_type, "s")

    def test_uniform_quantity_and_optional_fields(self):
        batch = self.prepare(
            [
                ["姓名", "手机号", "性别", "年龄", "职业"],
                ["合成客户甲", "13900000101", "女", 0, "教师"],
            ],
            mapping={"name": 0, "phone": 1, "gender": 2, "age": 3, "occupation": 4},
            uniform=3,
        )
        self.assertEqual(batch.total_cards, 3)
        row = batch.rows.get()
        self.assertEqual(row.normalized["age"], 0)
        self.assertEqual(row.normalized["gender"], "female")
        confirmed = imports.confirm_import(
            self.resource, batch.id, mapping_digest=batch.mapping_digest, version=batch.version
        )
        saved = imports.save_format(self.resource, batch.id, name="合成机构格式")
        encoded = json.dumps(saved.structure)
        self.assertNotIn("13900000101", encoded)
        self.assertNotIn("合成客户甲", encoded)
        self.assertNotIn("uniform_quantity", encoded)
        self.assertNotIn("教师", encoded)
        self.assertEqual(confirmed.status, "confirmed")

    def test_changed_mapping_revokes_confirmation_and_old_worker_cannot_overwrite(self):
        batch = self.prepare(
            [["姓名", "手机号", "数量", "客户编号"], ["合成甲", "13900000101", 1, "01"]]
        )
        old_version = batch.version
        batch = imports.confirm_import(
            self.resource, batch.id, mapping_digest=batch.mapping_digest, version=batch.version
        )
        batch = imports.configure_import(
            self.resource,
            batch.id,
            sheet="客户表",
            header_row=1,
            mapping={"name": 0, "phone": 1},
            quantity_mode="uniform",
            uniform_quantity=2,
            version=batch.version,
        )
        imports.validate_import(batch.id, old_version)
        batch.refresh_from_db()
        self.assertEqual(batch.status, "validating")
        self.assertIsNone(batch.confirmed_at)
        jobs.run_one()
        batch.refresh_from_db()
        self.assertEqual(batch.total_cards, 2)

    def test_aliases_ambiguity_and_org_isolation(self):
        suggestion = imports.suggest_mapping(["姓名", "手机号", "联系电话", "数量"])
        self.assertNotIn("phone", suggestion["mapping"])
        self.assertEqual(suggestion["ambiguous"]["phone"], [1, 2])
        self.assertTrue(suggestion["requires_confirmation"])
        batch = self.prepare(
            [["姓名", "手机号", "数量", "客户编号"], ["合成甲", "13900000101", 1, "01"]]
        )
        with self.assertRaises(BusinessError):
            imports.get_import(self.other, batch.id)
        self.assertEqual(api_client(self.other).get(f"/api/v1/imports/{batch.id}").status_code, 404)


class SalesConcurrencyTests(TransactionTestCase):
    def setUp(self):
        sales_setup(self)

    def test_two_orders_same_customer_resolve_once(self):
        first = order_fixture(self)
        second = order_fixture(self, actor=self.other, brand=self.other_brand)

        def run(order_id):
            close_old_connections()
            try:
                return sales.approve_order(
                    self.platform, order_id, approved=True, version=1, reason="并发合成审核"
                ).status
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as pool:
            self.assertEqual(list(pool.map(run, [first.id, second.id])), ["issued", "issued"])
        self.assertEqual(Customer.objects.count(), 1)
        self.assertEqual(CustomerSource.objects.count(), 2)
        self.assertEqual(Card.objects.count(), 4)

    def test_claim_against_cancel_never_partially_voids_claimed_batch(self):
        order = approve(self, order_fixture(self))
        customer = customers.register_verified_customer("13900000101")
        card = order.cards.first()

        def claim():
            close_old_connections()
            try:
                customers.claim(customer.id, card.id)
                return "claimed"
            except BusinessError:
                return "claim_blocked"
            finally:
                close_old_connections()

        def cancel():
            close_old_connections()
            try:
                item = sales.request_stop(
                    self.resource,
                    order.id,
                    action="cancel",
                    version=order.version,
                    reason="并发取消",
                )
                sales.review_stop(
                    self.platform, order.id, approved=True, version=item.version, reason="并发审核"
                )
                return "cancelled"
            except BusinessError:
                return "cancel_blocked"
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as pool:
            a, b = pool.submit(claim), pool.submit(cancel)
            results = [a.result(), b.result()]
        self.assertIn(results, [["claimed", "cancel_blocked"], ["claim_blocked", "cancelled"]])
        self.assertFalse(
            Card.objects.filter(activated_at__isnull=False, voided_at__isnull=False).exists()
        )
