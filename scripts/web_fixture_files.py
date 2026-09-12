"""Generate only synthetic upload fixtures inside the current server test report."""

import socket
import sys
from pathlib import Path

from openpyxl import Workbook
from PIL import Image, ImageDraw
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


if socket.gethostname() != 'VM-0-12-ubuntu':
    raise SystemExit('Server-only synthetic test fixtures')
directory = Path(sys.argv[1]).resolve()
root = Path('/home/ubuntu/ChiHuiTong/test-results')
if directory.parent != root or not directory.is_dir():
    raise SystemExit('Expected one existing isolated test report directory')
book = Workbook()
sheet = book.active
sheet.title = '合成客户'
sheet.append(['客户姓名', '手机号', '客户手机号', '开卡数量', '客户编号', '性别', '年龄', '职业'])
sheet.append(['合成批量一', '13800000801', '000', 2, 'EXTERNAL-801', '女', 29, '合成职业'])
sheet.append(['合成批量二', '+86 13800000802', '000', 3, 'EXTERNAL-802', '男', 31, '合成职业'])
book.save(directory/'synthetic-customers.xlsx')
# REQ-046 samples: non-first sheet, header at the supported boundary, >10 rows.
automatic = Workbook()
automatic.active.title = '说明'
automatic.active.append(['合成文件说明'])
customers = automatic.create_sheet('自动客户')
for _ in range(19):
    customers.append(['合成说明行'])
customers.append(['客户 姓名：', '手机号码', '购卡张数', '职业'])
for index in range(12):
    customers.append([f'自动客户{index+1:02d}', f'13800007{index:03d}', 1, '合成职业'])
automatic.save(directory/'synthetic-auto.xlsx')
missing = Workbook()
missing.active.append(['客户姓名', '联络号码', '备注'])
missing.active.append(['合成待对应客户', '13800000601', '不是数量'])
missing.save(directory/'synthetic-missing.xlsx')
invalid = Workbook()
invalid.active.append(['客户姓名', '手机号码', '开卡数量'])
invalid.active.append(['合成错误客户', '不是手机号', 0])
invalid.save(directory/'synthetic-invalid.xlsx')
# Template-style streaming XLSX legitimately omits optional dimension metadata.
template = Workbook(write_only=True)
sheet = template.create_sheet('客户清单')
sheet.append(['客户姓名', '手机号', '开卡数量', '资源方客户编号'])
sheet.append(['合成模板甲', '13800000501', 2, 'SOURCE-501'])
sheet.append(['合成模板乙', '13800000502', 3, 'SOURCE-502'])
template.save(directory/'synthetic-no-dimension.xlsx')
image = Image.new('RGB', (640, 360), 'white')
ImageDraw.Draw(image).text((30, 100), 'SYNTHETIC TEST ONLY - NOT A REAL PAYMENT / LICENSE', fill='black')
image.save(directory/'synthetic-proof.png')
pdf = PdfWriter()
page = pdf.add_blank_page(width=400, height=300)
font = DictionaryObject({NameObject('/Type'):NameObject('/Font'), NameObject('/Subtype'):NameObject('/Type1'), NameObject('/BaseFont'):NameObject('/Helvetica')})
page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):pdf._add_object(font)})})
stream = DecodedStreamObject()
stream.set_data(b'BT /F1 18 Tf 30 230 Td (SYNTHETIC TEST ONLY) Tj ET')
page[NameObject('/Contents')] = pdf._add_object(stream)
pdf.write(directory/'synthetic-static.pdf')
