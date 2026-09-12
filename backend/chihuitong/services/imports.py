import hashlib
import io
import json
import unicodedata

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from chihuitong.errors import BusinessError, require
from chihuitong.models import ImportBatch, ImportFormat, ImportRow, Outbox, SalesOrder

from .common import RESOURCE_KINDS, advance, audit, check_version
from .customers import match_preview, normalize_customer
from .files import checked_xlsx, file_bytes, validate_attachment_ids

ALIASES = {
    "name": [
        "姓名",
        "客户姓名",
        "客户名称",
        "持卡人姓名",
        "会员姓名",
        "用户姓名",
        "name",
        "customer_name",
    ],
    "phone": [
        "手机号",
        "手机号码",
        "客户手机号",
        "客户手机号码",
        "客户手机",
        "联系电话",
        "联系手机",
        "手机",
        "phone",
        "mobile",
        "phone_number",
    ],
    "quantity": [
        "开卡数量",
        "开卡张数",
        "购卡数量",
        "购卡张数",
        "采购数量",
        "数量",
        "张数",
        "卡数量",
        "quantity",
    ],
    "resource_customer_no": [
        "资源方客户编号",
        "客户编号",
        "客户编码",
        "会员编号",
        "客户号",
        "customerno",
    ],
    "gender": ["性别", "客户性别", "gender", "sex"],
    "age": ["年龄", "客户年龄", "age"],
    "occupation": ["职业", "客户职业", "occupation"],
}


def header_key(value):
    value = unicodedata.normalize("NFKC", str(value or "")).lower()
    # Normalize presentation punctuation, not arbitrary substrings (e.g. salesman phone).
    return "".join(char for char in value if char.isalnum())


def visible_imports(actor):
    qs = ImportBatch.objects.select_related("asset", "created_by")
    if actor.platform:
        return qs
    if actor.organization.kind not in RESOURCE_KINDS:
        return qs.none()
    qs = qs.filter(organization=actor.organization)
    return qs if actor.membership.role == "admin" else qs.filter(created_by=actor.membership)


def get_import(actor, batch_id, *, lock=False):
    qs = visible_imports(actor)
    if lock:
        qs = qs.select_for_update(of=("self",))
    batch = qs.filter(pk=batch_id).first()
    require(batch, "not_found", "导入记录不存在或无权访问", 404)
    return batch


@transaction.atomic
def create_import(actor, asset_id):
    require(actor.organization.kind in RESOURCE_KINDS, "forbidden", "仅资源方可导入客户名单", 403)
    asset = validate_attachment_ids(actor, [str(asset_id)], purposes={"sales_excel"})[0]
    batch = ImportBatch.objects.create(
        organization=actor.organization, created_by=actor.membership, asset=asset, status="queued"
    )
    Outbox.objects.create(
        kind="excel.inspect",
        dedup_key=f"excel.inspect:{batch.id}:1",
        payload={"batch_id": str(batch.id), "version": batch.version},
        available_at=timezone.now(),
    )
    audit(actor, batch, "import.uploaded")
    return batch


def open_book(batch):
    data = file_bytes(batch.asset)
    checked_xlsx(data)
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        require(0 < len(book.sheetnames) <= 20, "sheet_limit", "工作表过多，请拆分上传", 400)
        return book
    except (InvalidFileException, ValueError, KeyError) as exc:
        raise BusinessError("invalid_excel", "工作簿无法读取，请重新导出xlsx", 400) from exc


def safe_cell(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else str(value)
    return str(value)


def inspect_import(batch_id, expected_version):
    batch = ImportBatch.objects.select_related("asset").get(pk=batch_id)
    if batch.version != expected_version or batch.status != "queued":
        return
    book = open_book(batch)
    sheets, preview = [], {}
    try:
        for sheet in book:
            require(
                (sheet.max_row or 0) <= 10020 and (sheet.max_column or 0) <= 200,
                "excel_limit",
                "单表超过10000条数据或200列，请拆分上传",
                400,
            )
            sheets.append(
                {"name": sheet.title, "rows": sheet.max_row or 0, "columns": sheet.max_column or 0}
            )
            preview[sheet.title] = [
                [safe_cell(cell) for cell in row]
                for row in sheet.iter_rows(max_row=min(30, sheet.max_row or 0))
            ]
    finally:
        book.close()
    with transaction.atomic():
        batch = ImportBatch.objects.select_for_update().get(pk=batch_id)
        if batch.version != expected_version:
            return
        batch.sheets, batch.preview, batch.status = sheets, preview, "mapping"
        advance(batch, "sheets", "preview", "status")


@transaction.atomic
def configure_import(
    actor, batch_id, *, sheet, header_row, mapping, quantity_mode, uniform_quantity=None, version
):
    batch = get_import(actor, batch_id, lock=True)
    require(not actor.platform, "forbidden", "导入列映射由资源方确认", 403)
    check_version(batch, version)
    require(
        not SalesOrder.objects.filter(import_batch=batch).exists(),
        "import_consumed",
        "已提交订单的导入依据不可修改",
    )
    require(
        batch.status in {"mapping", "validated", "confirmed", "failed", "validating"},
        "invalid_state",
        "请等待工作簿读取完成",
    )
    sheet_meta = next((item for item in batch.sheets if item["name"] == sheet), None)
    require(
        sheet_meta and type(header_row) is int and 1 <= header_row <= min(20, sheet_meta["rows"]),
        "invalid_header",
        "请选择工作表和前20行内的表头",
        400,
    )
    require(
        isinstance(mapping, dict)
        and not (set(mapping) - set(ALIASES))
        and {"name", "phone"} <= set(mapping),
        "invalid_mapping",
        "请对应客户姓名和手机号列",
        400,
    )
    indexes = list(mapping.values())
    require(
        all(type(i) is int and 0 <= i < sheet_meta["columns"] for i in indexes)
        and len(indexes) == len(set(indexes)),
        "duplicate_mapping",
        "不同字段不能对应同一列或不存在的列",
        400,
    )
    require(
        quantity_mode in {"column", "uniform"},
        "quantity_required",
        "请选择按列读取或明确统一数量",
        400,
    )
    if quantity_mode == "column":
        require(
            "quantity" in mapping and uniform_quantity is None,
            "quantity_required",
            "按列读取须对应开卡数量，不叠加统一值",
            400,
        )
    else:
        require(
            type(uniform_quantity) is int
            and 1 <= uniform_quantity <= 100000
            and "quantity" not in mapping,
            "quantity_required",
            "请明确填写每人正整数张数，不同时选择数量列",
            400,
        )
    batch.configuration = {
        "sheet": sheet,
        "header_row": header_row,
        "mapping": mapping,
        "quantity_mode": quantity_mode,
        "uniform_quantity": uniform_quantity,
    }
    batch.mapping_digest = hashlib.sha256(
        json.dumps({"sha256": batch.asset.sha256, **batch.configuration}, sort_keys=True).encode()
    ).hexdigest()
    batch.confirmed_at, batch.status = None, "validating"
    batch.processed_rows = batch.total_rows = batch.error_rows = batch.total_cards = 0
    batch.failure_code = ""
    batch.rows.all().delete()
    advance(
        batch,
        "configuration",
        "mapping_digest",
        "confirmed_at",
        "status",
        "processed_rows",
        "total_rows",
        "error_rows",
        "total_cards",
        "failure_code",
    )
    Outbox.objects.create(
        kind="excel.validate",
        dedup_key=f"excel.validate:{batch.id}:{batch.version}",
        payload={"batch_id": str(batch.id), "version": batch.version},
        available_at=timezone.now(),
    )
    audit(actor, batch, "import.mapping_changed", version=batch.version)
    return batch


def validate_import(batch_id, expected_version):
    batch = ImportBatch.objects.select_related("asset").get(pk=batch_id)
    if batch.version != expected_version or batch.status != "validating":
        return
    book = open_book(batch)
    config = batch.configuration
    values, seen, duplicate_phones = [], {}, set()
    try:
        sheet = book[config["sheet"]]
        for number, cells in enumerate(
            sheet.iter_rows(min_row=config["header_row"] + 1), start=config["header_row"] + 1
        ):
            if all(cell.value in (None, "") for cell in cells):
                continue
            require(len(values) < 10000, "excel_limit", "单次数据超过10000行，请分批继续", 400)
            raw = {field: safe_cell(cells[index]) for field, index in config["mapping"].items()}
            if config["quantity_mode"] == "uniform":
                raw["quantity"] = config["uniform_quantity"]
            errors, normalized = [], {}
            try:
                require(
                    not any(cells[index].data_type == "f" for index in config["mapping"].values()),
                    "formula_field",
                    "关键字段包含公式，请转换为值后上传",
                    400,
                )
                normalized = normalize_customer(raw)
                match_preview(normalized)
                if normalized["phone"] in seen:
                    duplicate_phones.add(normalized["phone"])
                seen[normalized["phone"]] = number
            except BusinessError as exc:
                errors.append({"code": exc.code, "message": exc.message})
            values.append(
                ImportRow(
                    batch=batch,
                    row_number=number,
                    raw=raw,
                    normalized=normalized,
                    errors=errors,
                    status="invalid" if errors else "valid",
                )
            )
    finally:
        book.close()
    for row in values:
        if row.normalized.get("phone") in duplicate_phones:
            row.errors.append(
                {"code": "duplicate_phone", "message": "同一名单存在重复手机号，请核对"}
            )
            row.status = "invalid"
    require(values, "empty_import", "没有可导入的客户明细", 400)
    with transaction.atomic():
        batch = ImportBatch.objects.select_for_update().get(pk=batch_id)
        if batch.version != expected_version or batch.status != "validating":
            return
        batch.rows.all().delete()
        ImportRow.objects.bulk_create(values, batch_size=500)
        batch.total_rows = batch.processed_rows = len(values)
        batch.error_rows = sum(row.status == "invalid" for row in values)
        batch.total_cards = sum(
            row.normalized.get("quantity", 0) for row in values if row.status == "valid"
        )
        batch.status = "validated"
        advance(batch, "total_rows", "processed_rows", "error_rows", "total_cards", "status")


@transaction.atomic
def confirm_import(actor, batch_id, *, mapping_digest, version):
    batch = get_import(actor, batch_id, lock=True)
    require(not actor.platform, "forbidden", "由资源方确认列对应关系", 403)
    check_version(batch, version)
    require(
        batch.status == "validated"
        and batch.error_rows == 0
        and batch.total_rows > 0
        and 0 < batch.total_cards <= 100000,
        "import_errors",
        "请处理全部异常，单次开卡量不超过10万张，可分批继续",
    )
    require(mapping_digest == batch.mapping_digest, "stale_mapping", "列对应关系已变化，请重新确认")
    batch.status, batch.confirmed_at = "confirmed", timezone.now()
    advance(batch, "status", "confirmed_at")
    audit(actor, batch, "import.mapping_confirmed", version=batch.version)
    return batch


@transaction.atomic
def save_format(actor, batch_id, *, name):
    batch = get_import(actor, batch_id, lock=True)
    require(
        actor.organization.kind in RESOURCE_KINDS and batch.status == "confirmed",
        "unconfirmed_format",
        "请先确认本机构列对应关系",
        400,
    )
    require(
        isinstance(name, str) and 0 < len(name.strip()) <= 100,
        "invalid_name",
        "请填写格式名称",
        400,
    )
    config = batch.configuration
    headers = batch.preview[config["sheet"]][config["header_row"] - 1]
    normalized_headers = [header_key(value) for value in headers]
    field_headers = {field: normalized_headers[index] for field, index in config["mapping"].items()}
    structure = {
        "headers": normalized_headers,
        "field_headers": field_headers,
        "header_row": config["header_row"],
        "quantity_mode": config["quantity_mode"],
    }
    item = ImportFormat.objects.create(
        organization=actor.organization, name=name.strip(), structure=structure
    )
    audit(actor, batch, "import.format_saved", format_id=str(item.id))
    return item


def suggest_mapping(headers, stored_format=None):
    keys = [header_key(value) for value in headers]
    aliases = {field: {header_key(alias) for alias in values} for field, values in ALIASES.items()}
    suggestions, ambiguous = {}, {}
    if stored_format:
        require(
            sorted(keys)
            == sorted(header_key(value) for value in stored_format.structure["headers"]),
            "format_changed",
            "表头结构已变化，请重新对应列",
            400,
        )
        aliases = {
            field: {header_key(header)}
            for field, header in stored_format.structure["field_headers"].items()
        }
    for field, names in aliases.items():
        candidates = [i for i, key in enumerate(keys) if key and key in names]
        if len(candidates) == 1:
            suggestions[field] = candidates[0]
        elif len(candidates) > 1:
            ambiguous[field] = candidates
    return {"mapping": suggestions, "ambiguous": ambiguous, "requires_confirmation": True}


def recommend_import(batch):
    """Bounded, deterministic header suggestions; never infer customer fields from values."""
    formats = {}
    for item in ImportFormat.objects.filter(organization=batch.organization).order_by(
        "created_at", "id"
    ):
        structure = tuple(sorted(header_key(v) for v in item.structure["headers"]))
        formats.setdefault(structure, []).append(item)
    candidates = []
    for sheet in batch.sheets:
        for number, headers in enumerate(batch.preview.get(sheet["name"], [])[:20], 1):
            result = suggest_mapping(headers)
            saved = formats.get(tuple(sorted(header_key(v) for v in headers)), [])
            if saved:
                matches = [suggest_mapping(headers, item) for item in saved]
                # Conflicting historical formats must not silently choose one identity column.
                for field in ALIASES:
                    indexes = set()
                    for match in matches:
                        if field in match["mapping"]:
                            indexes.add(match["mapping"][field])
                        indexes.update(match["ambiguous"].get(field, []))
                    if indexes:
                        result["mapping"].pop(field, None)
                        result["ambiguous"].pop(field, None)
                        if len(indexes) == 1:
                            result["mapping"][field] = indexes.pop()
                        else:
                            result["ambiguous"][field] = sorted(indexes)
            known = set(result["mapping"]) | set(result["ambiguous"])
            score = len(known & {"name", "phone"}) * 100 + ("quantity" in known) * 20 + len(known)
            if score:
                candidates.append(
                    {
                        "sheet": sheet["name"],
                        "header_row": number,
                        "score": score,
                        "saved_format": bool(saved),
                        **result,
                    }
                )
    if not candidates:
        first = next((s for s in batch.sheets if s["rows"]), None)
        return {
            "sheet": first["name"] if first else "",
            "header_row": 1,
            "mapping": {},
            "ambiguous": {},
            "alternatives": 0,
            "saved_format": False,
        }
    candidates.sort(key=lambda item: item["score"], reverse=True)
    best = candidates[0]
    return {**best, "alternatives": sum(c["score"] == best["score"] for c in candidates) - 1}
